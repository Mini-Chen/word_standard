# app/word_generator.py
from docx import Document  # 读取创建word文件
import os  # 路径拼接
from datetime import datetime


def generate_word_file(data: dict) -> str:
    """
    data 结构示例:
    {
        "标题": "关于开展XX工作的通知",
        "主送单位": "各部门",
        "正文": "……正文内容……",
        "发文单位": "AI办公室",
        "发文日期": "2025年7月19日"
    }
    """
    template_path = os.path.join("templates", "redhead_template.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"模板不存在: {template_path}")

    doc = Document(template_path)

    # 简单占位符替换
    placeholders = {
        "{{标题}}": data.get("标题", ""),
        "{{主送单位}}": data.get("主送单位", ""),
        "{{正文内容}}": data.get("正文", data.get("正文内容", "")),
        "{{发文单位}}": data.get("发文单位", ""),
        "{{发文日期}}": data.get("发文日期", data.get("日期", datetime.now().strftime("%Y年%m月%d日")))
    }

    for p in doc.paragraphs:
        for k, v in placeholders.items():
            if k in p.text:
                p.text = p.text.replace(k, v)

    output_path = "输出_红头文件.docx"
    doc.save(output_path)
    return output_path
